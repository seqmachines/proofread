"""M1 setup writes: import local LibStructBench protocol folders into Atlas."""

import argparse
from concurrent.futures import ThreadPoolExecutor
import hashlib
import json
import re
import shutil
import subprocess
import tomllib
from urllib.parse import quote
from urllib.request import urlopen
import xml.etree.ElementTree as ET
from collections import Counter
from pathlib import Path
from zipfile import ZipFile

import pymupdf
from openpyxl import load_workbook
from pymongo import ReplaceOne

if __package__:
    from .db import get_db
else:
    from db import get_db

SEED_DIR = Path(__file__).with_name("seed")
BENCHMARK = "LibStructBench (Poon et al., 2026), protocol text and ground truth"
BASES = re.compile(r"(?<![A-Za-z])[ACGTUNacgtun]{10,}(?![A-Za-z])")
OPERATIONS = {"reverse_transcription", "template_switching", "pcr", "fragmentation", "ligation", "tagmentation", "other"}
PROTOCOLS = [
    ("smart_seq2", "Smart-seq2", "template_switching", "dev"),
    ("smart_seq3", "Smart-seq3", "template_switching", "dev"),
    ("cel_seq2", "CEL-seq2", "ivt", "dev"),
    ("sci_rna_seq", "sci-RNA-seq", "combinatorial_indexing", "regression"),
    ("plate_scatac_seq", "Plate scATAC-seq", "atac", "regression"),
    ("smart_seq3xpress", "Smart-seq3xpress", "template_switching", "transfer"),
]


def prepare(source_root):
    """Copy local source bundles and native annotations without editing them."""
    for slug, name, family, role in PROTOCOLS:
        destination = SEED_DIR / slug
        destination.mkdir(parents=True, exist_ok=True)
        (destination / "meta.json").write_text(json.dumps({
            "id": slug, "name": name, "family": family, "role": role,
        }, indent=2) + "\n")
        shutil.copyfile(source_root / "ground_truth" / slug / "groundtruth_library_generation_workflow.json",
                        destination / "ground_truth.json")
        shutil.copytree(source_root / "protocols" / slug, destination / "source", dirs_exist_ok=True)


def extract(source):
    """Yield source pages and bounded table sections with stable locators."""
    if source.suffix.lower() == ".pdf":
        with pymupdf.open(source) as document:
            for number, page in enumerate(document, 1):
                yield f"p{number:03}", number, "page", page.get_text(sort=True)
    elif source.suffix.lower() == ".xlsx":
        book = load_workbook(source, read_only=True, data_only=True)
        try:
            for number, sheet in enumerate(book, 1):
                columns = sheet.max_column or 1
                header = f"# {sheet.title}\n\n" + "\n".join([
                    "| " + " | ".join(f"Column {n + 1}" for n in range(columns)) + " |",
                    "| " + " | ".join("---" for _ in range(columns)) + " |"])
                rows, size, first, start, end = [], 0, True, 1, 0
                for row_number, row in enumerate(sheet.iter_rows(values_only=True), 1):
                    if not any(v is not None for v in row):
                        continue
                    cells = ["" if v is None else str(v).replace("|", "\\|").replace("\n", "<br>") for v in row]
                    line = "| " + " | ".join(cells) + " |"
                    if rows and (size + len(line) > 250_000 or len(rows) >= 1000):
                        locator = f"sheet{number:02}" if first else f"sheet{number:02}-rows{start:05}-{end:05}"
                        yield locator, number, "table", header + "\n" + "\n".join(rows)
                        rows, size, first = [], 0, False
                    if not rows:
                        start = row_number
                    rows.append(line)
                    size += len(line)
                    end = row_number
                if rows:
                    locator = f"sheet{number:02}" if first else f"sheet{number:02}-rows{start:05}-{end:05}"
                    yield locator, number, "table", header + "\n" + "\n".join(rows)
        finally:
            book.close()
    elif source.suffix.lower() == ".docx":
        w = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
        with ZipFile(source) as archive:
            document = ET.fromstring(archive.read("word/document.xml"))
        paragraphs = ["".join(t.text or "" for t in p.iter(w + "t")).strip()
                      for p in document.iter(w + "p")]
        paragraphs = [p for p in paragraphs if p]
        for start in range(0, len(paragraphs), 10):
            # DOCX pagination depends on Word. `page` is a logical block here.
            number = start // 10 + 1
            yield f"paragraphs{start + 1:03}-{min(start + 10, len(paragraphs)):03}", number, "page", "\n\n".join(paragraphs[start:start + 10])
    elif source.suffix.lower() == ".doc":
        if shutil.which("antiword"):
            command = ["antiword", str(source)]
        elif shutil.which("textutil"):
            command = ["textutil", "-convert", "txt", "-stdout", str(source)]
        else:
            raise RuntimeError("Legacy Word sources require antiword or macOS textutil")
        text = subprocess.run(command, check=True, capture_output=True, text=True).stdout
        for number, block in enumerate(text.split("\f"), 1):
            yield f"block{number:03}", number, "page", block
    else:
        raise ValueError(f"Unsupported source format: {source}")


def extract_chunks(folder, protocol_id):
    text_dir = folder / "text"
    text_dir.mkdir(exist_ok=True)
    chunks = []
    for source in sorted((folder / "source").iterdir()):
        if not source.is_file() or source.name.startswith("."):
            continue
        for locator, page, kind, raw in extract(source):
            if not raw.strip():
                continue
            text = BASES.sub("[nucleotide sequence omitted]", raw).strip()
            text = re.sub(r"(?:r[ACGU]){2,}(?:\+[ACGU])?", "[modified nucleotide sequence omitted]", text)
            text = f"Source: {source.name}; {locator}\n\n{text}"
            suffix = ".md" if kind == "table" else ".txt"
            (text_dir / f"{source.name}__{locator}{suffix}").write_text(text + "\n")
            chunks.append({
                "_id": f"{protocol_id}:{source.name}:{locator}", "protocol_id": protocol_id,
                "page": page, "kind": kind, "text": text, "source_file": source.name,
            })
    if not chunks:
        raise ValueError(f"{folder}: no source text was extracted")
    return chunks


def symbolic_segment(segment, origin="source"):
    name = segment["role"]
    description = (name + " " + segment.get("placeholder", "")).lower()
    if "umi" in description or "unique molecular" in description:
        kind = "umi"
    elif "barcode" in description:
        kind = "barcode"
    elif "index" in description or re.search(r"\bi[57]\b", description):
        kind = "index"
    elif "poly(a)" in description or "polya" in description or "poly-da" in description:
        kind = "polyA"
    elif "tso" in description or "template-switch" in description:
        kind = "tso"
    elif "handle" in description or "anchor" in description or "promoter" in description:
        kind = "handle"
    elif "adapter" in description or "adaptor" in description or "mosaic" in description:
        kind = "adapter"
    elif "primer" in description or "priming" in description or "poly-dt" in description or "poly(dt)" in description:
        kind = "primer"
    elif any(term in description for term in ("mrna", "cdna", "gdna", "transcript", "genomic")):
        kind = "insert"
    else:
        kind = "other"
    return {"name": BASES.sub("[sequence omitted]", name), "type": kind, "origin": origin}


def structural(value):
    """Keep native graph structure without nucleotide fields or free-text notes."""
    keys = {"workflow_id", "states", "transitions", "initial_state_ids", "final_outputs",
            "state_id", "strand_architecture", "reference_strand_id", "strands",
            "strand_id", "molecule_type", "orientation", "segments", "segment_id",
            "role", "structural_role", "support_status", "paired_regions",
            "paired_region_id", "side_1", "side_2", "segment_ids", "relationship",
            "discontinuities", "discontinuity_id", "after_segment_id", "before_segment_id", "kind",
            "transition_id", "operation", "substrate_state_ids", "product_state_ids",
            "carried_forward_product_ids", "discarded_product_ids", "oligo_ids", "modality"}
    if isinstance(value, dict):
        return {k: structural(v) for k, v in value.items() if k in keys}
    if isinstance(value, list):
        return [structural(v) for v in value]
    return BASES.sub("[sequence omitted]", value) if isinstance(value, str) else value


def validate_structure(state):
    """Reject dangling native pairing/discontinuity references before projecting."""
    strands = {s["strand_id"]: {p["segment_id"] for p in s["segments"]}
               for s in state["strands"]}
    if len(strands) != len(state["strands"]) or state["reference_strand_id"] not in strands:
        raise ValueError(f"Invalid strand IDs in {state['state_id']}")
    for strand in state["strands"]:
        if len(strands[strand["strand_id"]]) != len(strand["segments"]):
            raise ValueError(f"Duplicate segment IDs in {state['state_id']}")
    for pairing in state.get("paired_regions", []):
        for side in (pairing["side_1"], pairing["side_2"]):
            if side["strand_id"] not in strands or not set(side["segment_ids"]) <= strands[side["strand_id"]]:
                raise ValueError(f"Dangling pairing in {state['state_id']}")
    for gap in state.get("discontinuities", []):
        if gap["strand_id"] not in strands or not {
                gap["after_segment_id"], gap["before_segment_id"]} <= strands[gap["strand_id"]]:
            raise ValueError(f"Dangling discontinuity in {state['state_id']}")


def convert_ground_truth(native, protocol_id, *, origin="source", preserve_structure=False):
    """Project native T3 graphs to §2.1; never copy nucleotide sequence fields."""
    if native["protocol_id"] != protocol_id:
        raise ValueError("Task 3 protocol_id does not match metadata")
    states, transitions, projection_warnings = [], [], []
    for workflow in native["workflows"]:
        for state in workflow["states"]:
            strands = state["strands"]
            if not strands or any(s["orientation"] != "5_to_3" for s in strands):
                raise ValueError(f"Cannot project strand architecture in {state['state_id']}")
            validate_structure(state)
            # Keep the RNA template above its cDNA; otherwise use the native
            # reference strand as top. Native strands are all written 5'→3'.
            top = (next((s for s in strands if s["molecule_type"] == "RNA"), None)
                   if len(strands) <= 2 and not preserve_structure else None)
            if top is None:
                top = next(s for s in strands if s["strand_id"] == state["reference_strand_id"])
            # Preview an actual pairing partner when one is declared. Mixed
            # populations need not have a partner; never imply a false duplex.
            partner_ids = []
            for pair in state.get("paired_regions", []):
                for a, b in ((pair["side_1"], pair["side_2"]), (pair["side_2"], pair["side_1"])):
                    if a["strand_id"] == top["strand_id"]:
                        partner_ids.append(b["strand_id"])
            bottom = next((s for s in strands if s is not top and s["strand_id"] in partner_ids), None)
            if bottom is None and len(strands) == 2 and state["strand_architecture"] != "mixed_population":
                bottom = next(s for s in strands if s is not top)
            states.append({
                "id": state["state_id"], "label": BASES.sub("[sequence omitted]", state["name"]),
                "strands": {
                    "top": [symbolic_segment(s, origin) for s in top["segments"]],
                    "bottom": [symbolic_segment(s, origin) for s in reversed(bottom["segments"])] if bottom else [],
                },
                "origin": origin, "evidence": [], "skill_call_id": None,
                "review_status": "unreviewed", "stale_since_revision": None,
            })
            if preserve_structure or len(strands) > 2 or state["strand_architecture"] == "mixed_population":
                states[-1]["benchmark_structure"] = structural(state)
        for transition in workflow["transitions"]:
            op = transition["operation"]
            if "template_switch" in transition["transition_id"]:
                op = "template_switching"
            elif op not in OPERATIONS:
                op = "other"
            carried = transition["carried_forward_product_ids"]
            discarded = transition["discarded_product_ids"]
            if set(carried) & set(discarded):
                raise ValueError(f"Cannot project carried/discarded products in {transition['transition_id']}")
            if set(carried) | set(discarded) != set(transition["product_state_ids"]):
                raise ValueError(f"Unclassified product in {transition['transition_id']}")
            if not carried:
                # §2.1 edges require a carried product. Keep this native
                # disposition in scoring_workflows without inventing one.
                projection_warnings.append(f"{transition['transition_id']}: no carried product; "
                                           "discard-only operation is retained in the native record, not drawn as a carried edge.")
                continue
            pairs = [(a, b) for a in transition["substrate_state_ids"] for b in carried]
            for number, (source, target) in enumerate(pairs, 1):
                transition_id = transition["transition_id"]
                if len(pairs) > 1:
                    transition_id += f":{number}"
                transitions.append({
                    "id": transition_id, "from": source, "to": target, "op": op,
                    "skill_call_id": None, "evidence": [],
                    # Native Task 3 links to Task 2 by stable oligo identifier;
                    # preserve that symbolic name rather than inventing a label.
                    "oligos": list(transition["oligo_ids"]), "discarded": list(discarded),
                })
    state_ids = {s["id"] for s in states}
    if len(state_ids) != len(states) or len({t["id"] for t in transitions}) != len(transitions):
        raise ValueError("Task 3 has duplicate state or transition IDs")
    if not states or not transitions or any(t["from"] not in state_ids or t["to"] not in state_ids for t in transitions):
        raise ValueError("Task 3 has missing states or invalid transition endpoints")
    if any(state_id not in state_ids for t in transitions for state_id in t["discarded"]):
        raise ValueError("Task 3 has an unknown discarded product")
    # Setup stores a sequence-free graph for the isolated runtime scorer.
    scoring = structural(native["workflows"])
    for original, workflow in zip(native["workflows"], scoring):
        # Local benchmark records predate the final_outputs schema field.
        if "final_outputs" not in workflow:
            workflow["final_outputs"] = [
                {"state_id": sid, "modality": original.get("modality", "unknown")}
                for sid in original["final_state_ids"]]
        for transition in workflow["transitions"]:
            if "template_switch" in transition["transition_id"]:
                transition["operation"] = "template_switching"
    return {"protocol_id": protocol_id, "states": states, "transitions": transitions,
            "scoring_workflows": scoring, "projection_warnings": projection_warnings}


def load_seed():
    metas = sorted(SEED_DIR.glob("*/meta.json"))
    if not metas:
        raise FileNotFoundError(f"No protocol folders found in {SEED_DIR}; add the local benchmark first")
    protocols, chunks, truths = [], [], []
    for path in metas:
        meta = json.loads(path.read_text())
        protocol_id = meta.get("id", path.parent.name)
        gt_path = path.with_name("ground_truth.json")
        protocols.append({
            "_id": protocol_id, "id": protocol_id,
            "name": meta["name"], "family": meta["family"], "role": meta["role"],
            "has_gt": gt_path.exists(),
            "source_note": BENCHMARK,
        })
        chunks.extend(extract_chunks(path.parent, protocol_id))
        if gt_path.exists():
            truths.append(convert_ground_truth(json.loads(gt_path.read_text()), protocol_id))
        elif meta["role"] != "transfer":
            raise ValueError(f"{path.parent}: dev/regression protocols require ground_truth.json")

    if len({p["id"] for p in protocols}) != len(protocols):
        raise ValueError("Protocol IDs must be unique")
    roles = Counter(p["role"] for p in protocols)
    if roles not in (Counter(dev=3, regression=2, transfer=1), Counter(dev=2, regression=1, transfer=1)):
        raise ValueError("M1 requires 3 dev/2 regression/1 transfer, or the 2/1/1 cut")
    dev_families = Counter(p["family"] for p in protocols if p["role"] == "dev")
    target_family, count = dev_families.most_common(1)[0]
    regression_families = [p["family"] for p in protocols if p["role"] == "regression"]
    if count != 2 or target_family in regression_families or len(set(regression_families)) != len(regression_families):
        raise ValueError("Use two same-family dev protocols and regression protocols from different families")
    return protocols, chunks, truths


def seed():
    protocols, chunks, truths = load_seed()
    db = get_db()
    db.protocols.bulk_write([ReplaceOne({"_id": p["_id"]}, p, upsert=True) for p in protocols])
    db.chunks.bulk_write([ReplaceOne({"_id": c["_id"]}, c, upsert=True) for c in chunks])
    # Setup-only writes explicitly authorized by the user. Runtime reads belong
    # exclusively to verifier.py, called from the gate or the GT-diff route.
    db.ground_truth.bulk_write([
        ReplaceOne({"protocol_id": gt["protocol_id"]}, gt, upsert=True) for gt in truths
    ])
    print(f"Seeded {len(protocols)} protocols, {len(chunks)} chunks, {len(truths)} ground truths.")
    print("Transfer protocol was seeded only; no reconstruction or retrieval was run on it.")


def seed_benchmark_chunks(task_root):
    """Seed pinned source pages for the saved 20-protocol panel, without task code."""
    root = Path(task_root).expanduser().resolve(strict=True)
    db = get_db()
    protocols = sorted(db.benchmark_records.distinct("protocol_id"))
    if len(protocols) != 20:
        raise ValueError("Expected the selected 20 benchmark protocols")

    def one(protocol_id):
        task = (root / protocol_id).resolve(strict=True)
        if not task.is_relative_to(root):
            raise ValueError("Task path escaped the frozen bundle root")
        metadata = tomllib.loads((task / "task.toml").read_text())["metadata"]
        manifest = json.loads((task / "input_manifest.json").read_text())
        canonical = json.dumps(manifest, sort_keys=True, separators=(",", ":")).encode()
        if (manifest["protocol_id"] != protocol_id or
                hashlib.sha256(canonical).hexdigest() != metadata["source_manifest_sha256"]):
            raise ValueError(f"Frozen source manifest mismatch for {protocol_id}")
        repo, revision = metadata["input_repo"], metadata["input_revision"]
        if not re.fullmatch(r"[a-f0-9]{40}", revision):
            raise ValueError("Frozen input revision must be a full commit hash")
        folder = SEED_DIR / ("frozen-" + protocol_id)
        sources = folder / "source"
        sources.mkdir(parents=True, exist_ok=True)
        source_meta = {}
        for entry in manifest["sources"]:
            source = (sources / entry["local_path"]).resolve()
            if not source.is_relative_to(sources.resolve()) or source.parent != sources.resolve():
                raise ValueError("Frozen source must be a file directly within its bundle")
            if not source.exists() or hashlib.sha256(source.read_bytes()).hexdigest() != entry["sha256"]:
                url = f"https://huggingface.co/datasets/{quote(repo, safe='/')}/resolve/{revision}/{quote(entry['path'], safe='/')}"
                with urlopen(url, timeout=120) as response:
                    data = response.read()
                if hashlib.sha256(data).hexdigest() != entry["sha256"]:
                    raise ValueError(f"Frozen source hash mismatch for {protocol_id}: {source.name}")
                source.write_bytes(data)
            source_meta[source.name] = {"repo": repo, "revision": revision,
                "path": entry["path"], "sha256": entry["sha256"],
                "manifest_sha256": metadata["source_manifest_sha256"]}
        # Extract only the manifest's files, including when a prior cache has extras.
        chunks = []
        for filename, provenance in source_meta.items():
            for locator, page, kind, raw in extract(sources / filename):
                if not raw.strip():
                    continue
                text = BASES.sub("[nucleotide sequence omitted]", raw).strip()
                text = re.sub(r"(?:r[ACGU]){2,}(?:\+[ACGU])?", "[modified nucleotide sequence omitted]", text)
                chunks.append({"_id": f"{protocol_id}:{filename}:{locator}", "protocol_id": protocol_id,
                    "page": page, "kind": kind, "source_file": filename,
                    "text": f"Source: {filename}; {locator}\n\n{text}", "provenance": provenance})
        if not chunks:
            raise ValueError(f"No source text extracted for {protocol_id}")
        db.chunks.bulk_write([ReplaceOne({"_id": c["_id"]}, c, upsert=True) for c in chunks])
        print(json.dumps({"protocol_id": protocol_id, "sources": len(source_meta), "chunks": len(chunks)}), flush=True)
        return len(source_meta), len(chunks)

    with ThreadPoolExecutor(max_workers=4) as pool:
        results = list(pool.map(one, protocols))
    return {"protocols": len(protocols), "sources": sum(r[0] for r in results), "chunks": sum(r[1] for r in results)}


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source-root", type=Path, help="Copy the approved six bundles from local protocols-test first")
    parser.add_argument("--ground-truth-only", action="store_true", help="Reconvert native Task 3 without re-extracting source text")
    parser.add_argument("--benchmark-chunks", type=Path, help="Frozen task root for the selected 20 imported protocols")
    args = parser.parse_args()
    if args.benchmark_chunks:
        print(json.dumps(seed_benchmark_chunks(args.benchmark_chunks)))
        raise SystemExit(0)
    if args.source_root:
        prepare(args.source_root)
    if args.ground_truth_only:
        truths = [convert_ground_truth(json.loads(path.read_text()), path.parent.name)
                  for path in sorted(SEED_DIR.glob("*/ground_truth.json"))]
        if not truths:
            raise FileNotFoundError("No local Task 3 records found")
        get_db().ground_truth.bulk_write([
            ReplaceOne({"protocol_id": gt["protocol_id"]}, gt, upsert=True) for gt in truths
        ])
        print(f"Updated {len(truths)} converted ground truths; source chunks unchanged.")
    else:
        seed()
